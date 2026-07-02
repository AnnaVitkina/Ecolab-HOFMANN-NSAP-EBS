"""
Build a shipment/cost matrix workbook from combined processing files.
"""

from __future__ import annotations

import re
import sys
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from accessorial_costs import AccessorialCostRow, collect_accessorial_costs, get_used_matrix_tabs
from project_paths import INPUT_DIR, OUTPUT_DIR, PROCESSING_DIR, ensure_workspace_dirs
from process_tabs import parse_selection

PROJECT_NAME = "HOFMANN_Ecolab"
CURRENCY = "EUR"

COST_NAME_ROW = 1
APPLY_IF_ROW = 2
RATE_BY_ROW = 3
BRACKET_ROW = 4
COLUMN_HEADER_ROW = 5
DATA_START_ROW = 6

HEADER_FILL = PatternFill("solid", fgColor="D9D9D9")
BOLD = Font(bold=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)

SHIPMENT_HEADERS = [
    "Origin Country",
    "Origin Postal Code",
    "Destination Country",
    "Destination Postal Code",
]

TOLL_BLOCK = {
    "title": "Toll fee",
    "apply_if": "",
    "rate_by": "Rate by: Weight/kg",
}

GROUPAGE_BLOCK = {
    "title": "Transport cost (Ecolab BI Groupage)",
    "apply_if": "",
    "rate_by": "Rate by: Weight/kg",
}

TRANSFER_BLOCK = {
    "title": "Transport cost (Ecolab BI Transfer)",
    "apply_if": "",
    "rate_by": "Rate by: Weight/kg",
}

RETURN_BLOCK = {
    "title": "Transport cost (Return)",
    "apply_if": "",
    "rate_by": "Rate by: Weight/kg",
}


@dataclass
class CostColumnSpec:
    bracket_label: str
    rate_unit: str
    source_key: str


@dataclass
class CostBlock:
    title: str
    apply_if: str
    rate_by: str
    columns: list[CostColumnSpec] = field(default_factory=list)


@dataclass
class MatrixRow:
    shipment: dict[str, Any]
    costs: dict[tuple[str, str, str], float | None] = field(default_factory=dict)


@dataclass
class BracketColumn:
    source_key: str
    bracket_label: str
    column_index: int


@dataclass
class ParsedRateSheet:
    combined_tab: str
    file_name: str
    tab_name: str
    role: str
    brackets: list[BracketColumn]
    rates_by_postal: dict[str, dict[str, float | None]]


def cell_text(value: Any) -> str:
    if pd.isna(value):
        return ""
    return str(value).strip()


def rate_value(value: Any) -> float | None:
    if pd.isna(value):
        return None
    text = cell_text(value)
    if not text or text.lower() == "on request":
        return None
    try:
        return float(str(value).replace(",", "."))
    except (TypeError, ValueError):
        return None


def normalize_postal(value: Any) -> str | None:
    text = cell_text(value)
    if not text:
        return None
    if text.endswith(".0") and text[:-2].isdigit():
        text = text[:-2]
    if text.isdigit():
        return text.zfill(2)
    return text


def build_tab_to_file_map() -> dict[str, str]:
    mapping: dict[str, str] = {}
    for path in sorted(INPUT_DIR.glob("*.xlsx")):
        if path.name.startswith("~$"):
            continue
        try:
            workbook = pd.ExcelFile(path)
        except Exception:
            continue
        for sheet_name in workbook.sheet_names:
            mapping[sheet_name] = path.name
    return mapping


def resolve_source_file_name(
    tab_name: str,
    file_stem: str,
    combined_path: Path,
    tab_to_file: dict[str, str],
) -> str:
    if tab_name in tab_to_file:
        return tab_to_file[tab_name]
    if file_stem:
        return f"{file_stem}.xlsx"
    return combined_path.name


def split_combined_tab_name(combined_tab: str) -> tuple[str, str]:
    if "__" in combined_tab:
        file_part, tab_part = combined_tab.split("__", 1)
        return file_part, tab_part
    return "", combined_tab


def classify_tab(combined_tab: str) -> str:
    _, tab_name = split_combined_tab_name(combined_tab)
    lowered = tab_name.lower()
    if "maut" in lowered:
        return "maut"
    if "groupage" in lowered:
        return "groupage"
    if "transfer" in lowered:
        return "transfer"
    return "other"


def ensure_directories() -> None:
    ensure_workspace_dirs()


def matrix_output_path(destination_country: str) -> Path:
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    filename = f"{PROJECT_NAME}_{destination_country}_rate_matrix_{timestamp}.xlsx"
    return OUTPUT_DIR / filename


def list_combined_workbooks() -> list[Path]:
    ensure_directories()
    files = sorted(PROCESSING_DIR.glob("combined_*.xlsx"), reverse=True)
    return [path for path in files if not path.name.startswith("~$")]


def prompt_combined_workbook() -> Path:
    files = list_combined_workbooks()
    if not files:
        print(f"No combined workbooks found in {PROCESSING_DIR}. Run process_tabs.py first.")
        sys.exit(1)

    print("\nCombined workbooks in processing/:")
    for index, path in enumerate(files, start=1):
        marker = " (latest)" if index == 1 else ""
        print(f"  {index}. {path.name}{marker}")
    print("\nPress Enter for latest, or enter file number:")

    while True:
        raw = input("> ").strip()
        if not raw:
            return files[0]
        if raw.isdigit():
            index = int(raw) - 1
            if 0 <= index < len(files):
                return files[index]
        print("Invalid selection. Try again.")


def prompt_destination_country(default: str = "DE") -> str:
    while True:
        raw = input(f"Destination Country [{default}]: ").strip().upper()
        if not raw:
            return default
        if re.fullmatch(r"[A-Z]{2}", raw):
            return raw
        print("Please enter a 2-letter country code (e.g. DE).")


def prompt_return_tab(default_tab: str, available_tabs: list[str]) -> str:
    print(f"\nTransport cost (Return) will use tab: {default_tab}")
    answer = input("Is this correct? [Y/n]: ").strip().lower()
    if answer in {"", "y", "yes"}:
        return default_tab

    print("\nChoose tab for Transport cost (Return):")
    for index, tab in enumerate(available_tabs, start=1):
        print(f"  {index}. {tab}")

    while True:
        raw = input("Enter tab number: ").strip()
        try:
            indices = parse_selection(raw, len(available_tabs))
            if len(indices) == 1:
                return available_tabs[indices[0]]
        except ValueError as exc:
            print(f"Invalid input: {exc}")
        print("Please enter exactly one tab number.")


def normalize_maut_bracket_label(label: str) -> str:
    text = cell_text(label)
    if not text:
        return text
    if text.upper() == "FTL":
        return "FTL"

    open_end_match = re.match(r"^(\d+)-$", text)
    if open_end_match:
        lower_bound = int(open_end_match.group(1))
        return f">{lower_bound - 1}"

    range_match = re.match(r"^(\d+)-(\d+)$", text)
    if range_match:
        upper_bound = int(range_match.group(2))
        return f"<={upper_bound}"

    return text


def drop_duplicate_ftl_bracket(
    brackets: list[BracketColumn],
    rates_by_postal: dict[str, dict[str, float | None]],
) -> list[BracketColumn]:
    ftl_brackets = [bracket for bracket in brackets if bracket.bracket_label == "FTL"]
    over_brackets = [bracket for bracket in brackets if bracket.bracket_label.startswith(">")]
    if not ftl_brackets or not over_brackets:
        return brackets

    ftl_bracket = ftl_brackets[-1]
    over_bracket = over_brackets[-1]
    values_match = True
    for rates in rates_by_postal.values():
        ftl_value = rates.get(ftl_bracket.source_key)
        over_value = rates.get(over_bracket.source_key)
        if ftl_value is None and over_value is None:
            continue
        if ftl_value != over_value:
            values_match = False
            break

    if not values_match:
        return brackets

    return [bracket for bracket in brackets if bracket.source_key != ftl_bracket.source_key]


def find_maut_brackets(header_row: pd.Series) -> list[BracketColumn]:
    brackets: list[BracketColumn] = []
    for index, value in enumerate(header_row):
        if index == 0:
            continue
        label = cell_text(value)
        if not label:
            continue
        brackets.append(
            BracketColumn(
                source_key=f"col_{index}",
                bracket_label=normalize_maut_bracket_label(label),
                column_index=index,
            )
        )
    return brackets


def numeric_bracket_label(upper_bound: float) -> str:
    if upper_bound == int(upper_bound):
        return f"<={int(upper_bound)}"
    return f"<={upper_bound}"


def find_numeric_brackets(header_row: pd.Series) -> list[BracketColumn]:
    numeric_columns: list[tuple[int, float]] = []
    for index, value in enumerate(header_row):
        if index == 0:
            continue
        if pd.isna(value):
            continue
        try:
            numeric_columns.append((index, float(value)))
        except (TypeError, ValueError):
            continue

    if not numeric_columns:
        return []

    numeric_columns.sort(key=lambda item: item[1])
    last_index = len(numeric_columns) - 1
    brackets: list[BracketColumn] = []
    for offset, (column_index, upper_bound) in enumerate(numeric_columns):
        if offset == last_index:
            upper_label = numeric_bracket_label(upper_bound)
            over_label = (
                f">{int(upper_bound)}"
                if upper_bound == int(upper_bound)
                else f">{upper_bound}"
            )
            brackets.append(
                BracketColumn(
                    source_key=f"col_{column_index}",
                    bracket_label=upper_label,
                    column_index=column_index,
                )
            )
            brackets.append(
                BracketColumn(
                    source_key=f"col_{column_index}_gt",
                    bracket_label=over_label,
                    column_index=column_index,
                )
            )
            continue

        brackets.append(
            BracketColumn(
                source_key=f"col_{column_index}",
                bracket_label=numeric_bracket_label(upper_bound),
                column_index=column_index,
            )
        )
    return brackets


def parse_rate_sheet(
    combined_path: Path,
    combined_tab: str,
    tab_to_file: dict[str, str],
) -> ParsedRateSheet | None:
    raw = pd.read_excel(combined_path, sheet_name=combined_tab, header=None)
    if raw.empty:
        return None

    role = classify_tab(combined_tab)
    file_stem, tab_name = split_combined_tab_name(combined_tab)
    file_name = resolve_source_file_name(tab_name, file_stem, combined_path, tab_to_file)

    if role == "maut":
        brackets = find_maut_brackets(raw.iloc[0])
    elif role in {"groupage", "transfer"}:
        brackets = find_numeric_brackets(raw.iloc[0])
    else:
        return None

    if not brackets:
        print(f"  Skipping '{combined_tab}': could not detect weight bracket columns.")
        return None

    rates_by_postal: dict[str, dict[str, float | None]] = {}
    for row_index in range(1, len(raw)):
        postal = normalize_postal(raw.iloc[row_index, 0])
        if not postal:
            continue
        row_rates: dict[str, float | None] = {}
        for bracket in brackets:
            row_rates[bracket.source_key] = rate_value(
                raw.iloc[row_index, bracket.column_index]
            )
        rates_by_postal[postal] = row_rates

    if role == "maut":
        brackets = drop_duplicate_ftl_bracket(brackets, rates_by_postal)

    return ParsedRateSheet(
        combined_tab=combined_tab,
        file_name=file_name,
        tab_name=tab_name,
        role=role,
        brackets=brackets,
        rates_by_postal=rates_by_postal,
    )


def make_cost_block(block_meta: dict[str, str], sheet: ParsedRateSheet | None) -> CostBlock | None:
    if sheet is None:
        return None
    columns = [
        CostColumnSpec(bracket.bracket_label, "Flat", bracket.source_key)
        for bracket in sheet.brackets
    ]
    return CostBlock(**block_meta, columns=columns)


def build_cost_blocks(
    parsed_sheets: dict[str, ParsedRateSheet],
    *,
    return_sheet: ParsedRateSheet,
) -> list[CostBlock]:
    blocks: list[CostBlock] = []
    for block_meta, role in (
        (TOLL_BLOCK, "maut"),
        (GROUPAGE_BLOCK, "groupage"),
        (TRANSFER_BLOCK, "transfer"),
    ):
        block = make_cost_block(block_meta, parsed_sheets.get(role))
        if block is not None:
            blocks.append(block)

    return_block = make_cost_block(RETURN_BLOCK, return_sheet)
    if return_block is not None:
        blocks.append(return_block)
    return blocks


def cost_key(block: CostBlock, spec: CostColumnSpec) -> tuple[str, str, str]:
    return (block.title, spec.bracket_label, spec.rate_unit)


def has_cost(value: float | None) -> bool:
    return value is not None and value != 0.0


def block_has_any_cost(matrix_row: MatrixRow, block: CostBlock) -> bool:
    return any(has_cost(matrix_row.costs.get(cost_key(block, spec))) for spec in block.columns)


def build_matrix_rows(
    parsed_sheets: dict[str, ParsedRateSheet],
    *,
    destination_country: str,
    return_tab: str,
) -> list[MatrixRow]:
    groupage = parsed_sheets.get("groupage")
    if groupage is None:
        raise ValueError("No Ecolab BI Groupage tab found in the combined workbook.")

    maut = parsed_sheets.get("maut")
    transfer = parsed_sheets.get("transfer")
    return_sheet = None
    for sheet in parsed_sheets.values():
        if sheet.combined_tab == return_tab:
            return_sheet = sheet
            break
    if return_sheet is None:
        return_sheet = groupage

    postal_codes = sorted(groupage.rates_by_postal)
    matrix_rows: list[MatrixRow] = []

    for postal in postal_codes:
        shipment = {
            "Origin Country": "",
            "Origin Postal Code": "",
            "Destination Country": destination_country,
            "Destination Postal Code": postal,
        }
        row = MatrixRow(shipment=shipment)

        def fill_block(block_meta: dict[str, str], sheet: ParsedRateSheet | None) -> None:
            if sheet is None:
                return
            rates = sheet.rates_by_postal.get(postal, {})
            block = CostBlock(**block_meta, columns=[])
            for bracket in sheet.brackets:
                spec = CostColumnSpec(bracket.bracket_label, "Flat", bracket.source_key)
                block.columns.append(spec)
                row.costs[cost_key(block, spec)] = rates.get(bracket.source_key)

        fill_block(TOLL_BLOCK, maut)
        fill_block(GROUPAGE_BLOCK, groupage)
        fill_block(TRANSFER_BLOCK, transfer)
        fill_block(RETURN_BLOCK, return_sheet or groupage)
        matrix_rows.append(row)

    return matrix_rows


def block_uses_shared_currency(block: CostBlock) -> bool:
    return block.rate_by == "Rate by: Weight/kg"


def block_column_width(block: CostBlock) -> int:
    if block_uses_shared_currency(block):
        return 1 + len(block.columns)
    return len(block.columns) * 2


def filter_cost_blocks_with_data(
    matrix_rows: list[MatrixRow],
    cost_blocks: list[CostBlock],
) -> list[CostBlock]:
    return [
        block
        for block in cost_blocks
        if any(block_has_any_cost(row, block) for row in matrix_rows)
    ]


def write_rates_sheet(
    worksheet,
    matrix_rows: list[MatrixRow],
    cost_blocks: list[CostBlock],
) -> None:
    shipment_count = len(SHIPMENT_HEADERS)

    def write_cost_header_row(row_index: int, values: list[str]) -> None:
        column_index = shipment_count + 1
        for block_index, block in enumerate(cost_blocks):
            block_width = block_column_width(block)
            value = values[block_index] if block_index < len(values) else ""
            cell = worksheet.cell(row=row_index, column=column_index, value=value)
            cell.font = BOLD
            cell.fill = HEADER_FILL
            cell.alignment = LEFT
            if block_width > 1:
                worksheet.merge_cells(
                    start_row=row_index,
                    start_column=column_index,
                    end_row=row_index,
                    end_column=column_index + block_width - 1,
                )
            column_index += block_width

    write_cost_header_row(COST_NAME_ROW, [block.title for block in cost_blocks])
    write_cost_header_row(APPLY_IF_ROW, [block.apply_if for block in cost_blocks])
    write_cost_header_row(RATE_BY_ROW, [block.rate_by for block in cost_blocks])

    for col_idx, header in enumerate(SHIPMENT_HEADERS, start=1):
        cell = worksheet.cell(row=COLUMN_HEADER_ROW, column=col_idx, value=header)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = LEFT

    column_index = shipment_count + 1
    for block in cost_blocks:
        currency_cell = worksheet.cell(row=COLUMN_HEADER_ROW, column=column_index, value=CURRENCY)
        currency_cell.font = BOLD
        currency_cell.fill = HEADER_FILL
        for spec_offset, spec in enumerate(block.columns):
            spec_col = column_index + 1 + spec_offset
            bracket_cell = worksheet.cell(row=BRACKET_ROW, column=spec_col, value=spec.bracket_label)
            bracket_cell.font = BOLD
            bracket_cell.fill = HEADER_FILL
            bracket_cell.alignment = CENTER
            unit_cell = worksheet.cell(row=COLUMN_HEADER_ROW, column=spec_col, value=spec.rate_unit)
            unit_cell.font = BOLD
            unit_cell.fill = HEADER_FILL
            unit_cell.alignment = CENTER
        column_index += block_column_width(block)

    for row_offset, matrix_row in enumerate(matrix_rows):
        excel_row = DATA_START_ROW + row_offset
        for col_idx, header in enumerate(SHIPMENT_HEADERS, start=1):
            worksheet.cell(row=excel_row, column=col_idx, value=matrix_row.shipment.get(header))

        column_index = shipment_count + 1
        for block in cost_blocks:
            block_has_costs = block_has_any_cost(matrix_row, block)
            if block_has_costs:
                worksheet.cell(row=excel_row, column=column_index, value=CURRENCY)
            for spec_offset, spec in enumerate(block.columns):
                spec_col = column_index + 1 + spec_offset
                value = matrix_row.costs.get(cost_key(block, spec))
                if has_cost(value):
                    value_cell = worksheet.cell(row=excel_row, column=spec_col, value=value)
                    value_cell.number_format = "0.00"
                elif block_has_costs:
                    value_cell = worksheet.cell(row=excel_row, column=spec_col, value=0)
                    value_cell.number_format = "0.00"
            column_index += block_column_width(block)

    for col_idx in range(1, worksheet.max_column + 1):
        worksheet.column_dimensions[get_column_letter(col_idx)].width = 18


def write_accessorial_sheet(
    worksheet,
    accessorial_rows: list[AccessorialCostRow],
) -> None:
    headers = ["Cost name", "Cost price", "Rate by", "Apply if"]
    for col_idx, header in enumerate(headers, start=1):
        cell = worksheet.cell(row=1, column=col_idx, value=header)
        cell.font = BOLD
        cell.fill = HEADER_FILL
        cell.alignment = LEFT

    for row_index, row in enumerate(accessorial_rows, start=2):
        worksheet.cell(row=row_index, column=1, value=row.cost_name)
        price_cell = worksheet.cell(row=row_index, column=2, value=row.cost_price)
        price_cell.number_format = "0.00"
        worksheet.cell(row=row_index, column=3, value=row.rate_by)
        worksheet.cell(row=row_index, column=4, value=row.apply_if)

    for col_idx in range(1, len(headers) + 1):
        worksheet.column_dimensions[get_column_letter(col_idx)].width = 24


def write_matrix_workbook(
    matrix_rows: list[MatrixRow],
    cost_blocks: list[CostBlock],
    output_path: Path,
    *,
    accessorial_rows: list[AccessorialCostRow] | None = None,
) -> Path:
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Rates"
    active_blocks = filter_cost_blocks_with_data(matrix_rows, cost_blocks)
    write_rates_sheet(worksheet, matrix_rows, active_blocks)

    accessorial_worksheet = workbook.create_sheet("Accessorial costs")
    write_accessorial_sheet(accessorial_worksheet, accessorial_rows or [])

    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)
    return output_path


def build_matrix_from_workbook(
    combined_path: Path,
    *,
    output_path: Path | None = None,
) -> Path:
    excel_file = pd.ExcelFile(combined_path)
    parsed_by_role: dict[str, ParsedRateSheet] = {}
    available_tabs = excel_file.sheet_names
    tab_to_file = build_tab_to_file_map()

    print(f"\nReading tabs from {combined_path.name}:")
    for combined_tab in available_tabs:
        parsed = parse_rate_sheet(combined_path, combined_tab, tab_to_file)
        if parsed is None:
            continue
        parsed_by_role[parsed.role] = parsed
        print(
            f"  {parsed.role}: {combined_tab} "
            f"({len(parsed.rates_by_postal)} postal rows, {len(parsed.brackets)} brackets)"
        )

    if "groupage" not in parsed_by_role:
        raise ValueError("Combined workbook must include an Ecolab BI Groupage tab.")

    destination_country = prompt_destination_country("DE")
    default_return_tab = parsed_by_role["groupage"].combined_tab
    return_tab = prompt_return_tab(default_return_tab, available_tabs)

    return_sheet = None
    for sheet in parsed_by_role.values():
        if sheet.combined_tab == return_tab:
            return_sheet = sheet
            break
    if return_sheet is None:
        return_sheet = parsed_by_role["groupage"]

    matrix_rows = build_matrix_rows(
        parsed_by_role,
        destination_country=destination_country,
        return_tab=return_tab,
    )
    if not matrix_rows:
        raise ValueError("No matrix rows were produced.")

    target_path = output_path or matrix_output_path(destination_country)
    cost_blocks = build_cost_blocks(parsed_by_role, return_sheet=return_sheet)

    used_tabs = get_used_matrix_tabs(parsed_by_role, return_tab)
    print("\nCollecting accessorial costs from unused tabs:")
    accessorial_rows = collect_accessorial_costs(excel_file, used_tabs)

    write_matrix_workbook(
        matrix_rows,
        cost_blocks,
        target_path,
        accessorial_rows=accessorial_rows,
    )
    return target_path


def run_matrix_build(
    combined_path: Path,
    *,
    output_path: Path | None = None,
) -> Path:
    print(f"\n--- Building matrix from: {combined_path.name} ---")
    result_path = build_matrix_from_workbook(combined_path, output_path=output_path)
    row_count = len(pd.read_excel(result_path, sheet_name="Rates", header=None)) - DATA_START_ROW + 1
    print(f"\nSaved matrix workbook: {result_path}")
    print(f"  Matrix rows: {row_count}")
    return result_path


def main() -> None:
    ensure_directories()
    combined_path = prompt_combined_workbook()
    run_matrix_build(combined_path)


if __name__ == "__main__":
    main()
